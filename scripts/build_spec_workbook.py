"""Builds docs/spec/RCM_Dashboard_Spec.xlsx from docs/spec/catalog.json (scripts/export-spec.ts)
and docs/spec/visuals.json (visual inventory read from the running prototype).
Scope: ELT Summary and Analytics dashboards. Worklists (v2) are out of scope.
Run: python3 scripts/build_spec_workbook.py
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

CAT = json.load(open('docs/spec/catalog.json'))
VIS = json.load(open('docs/spec/visuals.json'))
OUT = 'docs/spec/RCM_Dashboard_Spec.xlsx'

M = {m['id']: m for m in CAT['metrics']}
PAGE = {p['id']: p for p in CAT['pages']}
FIELD = CAT['fields']
UNIT = {'pct': 'Percent', 'days': 'Days', 'usd': 'Dollars', 'count': 'Count', 'sec': 'Seconds', 'ratio': 'Ratio'}
DIRN = {'up': 'Higher is better', 'down': 'Lower is better', 'none': 'No preferred direction'}
FW = {m['id']: (g['name'], m) for g in CAT['framework'] for m in g['metrics']}

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='1F3864')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='D9D9D9')

# ------------------------------------------------------------------ dashboards and views
DASHBOARDS = [
    ('DB-01', 'ELT Summary', 'Executive view of revenue cycle performance against target ("See": How are we doing?). Low density, high signal. '
     'Built from standard BI components: KPI cards, trend charts, a metric table with conditional formatting, ranked variance tables, slicers, tooltips and drill-through.',
     'Executive leadership team (CFO, VP Revenue Cycle)', 'Single page. Every metric drills to its Analytics view.',
     'Compact slicer bar: Period type, Period, Hospital (or region).'),
    ('DB-02', 'Analytics', 'Interactive analysis of the revenue cycle ("Understand": What is happening, where, and why?). Medium-high density. '
     'Tabs follow the revenue cycle; drill-through pages give metric, hospital and account detail.',
     'Revenue cycle directors, managers and analysts', 'Tabs across the top (one per view); drill-through views open from visuals; breadcrumbs and Back.',
     'Full filter pane (see Filters & Dimensions) plus page-specific filters.'),
]
ANALYTICS_ORDER = ['cycle', 'access', 'midcycle', 'billing', 'denials', 'ar', 'cash', 'payers', 'facilities', 'definitions', 'metric', 'facility', 'accounts']
VIEW_EXTRA = {
    'executive': 'Five sections: Key Performance Indicators, Performance Trends, Performance vs. Target, Areas of Focus, Performance Highlights.',
    'cycle': 'Full enterprise scorecard (P1-P13 first), ranked hospital exceptions, lifecycle stages, hospital performance matrix.',
    'access': 'Contact center, scheduling, pre-registration and financial clearance (72-hour windows), POS collections, eligibility and Medicaid enrollment, registration quality.',
    'midcycle': 'Charge lag, late charges, missed charges, final coded volume, coding turnaround by type, CDI, documentation denials.',
    'billing': 'DNFB $ and days, hold reasons, aging, claims exported, billed not exported, first pass yield, clean claims, edits, billing lag.',
    'denials': 'Decomposition tree, drivers, payer x category, appeal volume, outcome and status.',
    'ar': 'Aging trend, hospital, payer, financial class, status, credit balances, concentration, aging matrix, largest accounts.',
    'cash': 'Month pace, YTD by hospital, payer collections, non-charge transactions, retractions, contractuals, life-to-date yield by service month.',
    'payers': 'Mix, A/R days, denial rate, net collection, contractual adjustment, payment variance, underpayment flags.',
    'facilities': 'Sortable hospital comparison with status, trend and drill-through; ranking by a selectable metric.',
    'definitions': 'Reference tabs: Metric catalog, Client KPI coverage, Data sources, Dimensions, States, Data contract.',
    'metric': 'Generic drill-through for any catalog metric: trend with target / PY / rolling-3, breakdowns, contribution to change, peer comparison, small multiples.',
    'facility': 'One hospital vs target, peers and system, by stage.',
    'accounts': 'The accounts behind an aggregate (synthetic IDs, no patient data in the prototype).',
}
ENTRY = {'metric': 'Any KPI card, scorecard row or trend (drill-through)', 'facility': 'Any hospital link or bar in Drill mode',
         'accounts': 'Aggregate visuals on A/R, Billing / DNFB, Denials (drill-through)'}

views = [('V-01', 'DB-01', 'executive', 'ELT Summary', 'Page')]
for i, pid in enumerate(ANALYTICS_ORDER, start=2):
    views.append((f'V-{i:02d}', 'DB-02', pid, PAGE[pid]['title'], 'Drill-through' if pid in ENTRY else 'Tab'))
VIEW_BY_PAGE = {v[2]: v for v in views}

def page_filters(pid):
    if pid == 'executive':
        return 'Period type, Period, Hospital (or region)'
    base = 'Global filter pane'
    pf = PAGE[pid].get('pageFilters') or []
    return base + ('; page filters: ' + ', '.join(FIELD[f] for f in pf) if pf else '')

# ------------------------------------------------------------------ spec rows
rows = []  # dicts
visuals = []  # (vz_id, view_id, title, type, dims, interaction, notes)

def mname(mid):
    m = M[mid]
    return (f"{m['code']} " if m['code'] else '') + m['name']

def add_visual(view_id, title, vtype, dims, interaction, metric_roles, notes=''):
    n = sum(1 for v in visuals if v[1] == view_id) + 1
    vz = f'{view_id}-VZ{n:02d}'
    visuals.append((vz, view_id, title, vtype, dims, interaction, notes))
    if not metric_roles:
        rows.append(dict(vz=vz, metric=None, role='No metric (reference or detail list)', notes=notes))
    seen = set()
    for mid, role, note in metric_roles:
        if mid in seen:
            continue
        seen.add(mid)
        rows.append(dict(vz=vz, metric=mid, role=role, notes=note or ''))

# ELT Summary (hand-mapped from src/pages/EltSummary.tsx; ELT sections are not Visual components)
ELT_KPIS = ['gross_ar_days', 'cash_pct_npsr', 'denial_dollar_rate', 'dnfb_days', 'ar_gt90_pct']
ELT_TRENDS = ['gross_ar_days', 'cash_pct_npsr', 'denial_dollar_rate']
FW_ALL = [m['id'] for g in CAT['framework'] for m in g['metrics']]
add_visual('V-01', 'Key Performance Indicators', 'KPI cards (5): actual, target, variance, 3-month change, 12-month sparkline', '',
           'Card click drills to the metric\'s Analytics view (P1, P4 -> A/R; P5 -> Billing / DNFB; P8 -> Cash; P10 -> Denials)',
           [(i, 'KPI card', '') for i in ELT_KPIS])
for mid in ELT_TRENDS:
    add_visual('V-01', f'Performance Trends: {M[mid]["name"]}', 'Line chart, 12 months, with target line and off-target range', 'Month',
               'Click opens Metric Analysis (drill-through)', [(mid, 'Primary', '')])
add_visual('V-01', 'Performance vs. Target', 'Metric table with conditional formatting, grouped Balance Sheet / Income Statement: actual, target, variance, status, 3-month change, 12-month trend, benchmark',
           'Framework group', 'Row click drills to the metric\'s Analytics view; hospital link applies that hospital first',
           [(i, 'Row in table', '') for i in FW_ALL])
add_visual('V-01', 'Areas of Focus', 'Ranked variance table (up to 5 rows)', 'Metric, key driver hospital',
           'Row click drills to the Analytics view; hospital link filters to that hospital',
           [(i, 'Candidate row (shown when off target)', 'Rows chosen at run time: off-target metrics with largest unfavorable variance relative to target.') for i in FW_ALL])
add_visual('V-01', 'Performance Highlights', 'Ranked variance table (up to 3 rows)', 'Metric, leading hospital',
           'Row click drills to the Analytics view; hospital link filters to that hospital',
           [(i, 'Candidate row (shown when on target)', 'Rows chosen at run time: on-target metrics with largest favorable variance.') for i in FW_ALL])

SKIP_TITLES = {'Ranking metric'}  # a selector control, folded into the ranking chart
for pg in VIS['pages']:
    pid = pg['page']
    vid = VIEW_BY_PAGE[pid][0]
    generic = pid == 'metric'
    for k, strip in enumerate([s for s in pg['strips'] if s]):
        add_visual(vid, f'KPI cards ({k + 1})' if len(pg['strips']) > 1 else 'KPI cards', 'KPI card strip: value, target, variance, change, sparkline', '',
                   'Card click opens Metric Analysis (drill-through)', [(i, 'KPI card', '') for i in strip])
    for v in pg['visuals']:
        if v['title'] in SKIP_TITLES:
            continue
        spec = v['spec'] or {}
        title = v['pid'] if generic or v['title'] != v['pid'] else v['title']
        roles = []
        if generic:
            roles = []
        else:
            prim = v['metric'] or None
            ms = spec.get('metrics') or []
            if prim:
                roles.append((prim, 'Primary', ''))
            for i in ms:
                t = (spec.get('type') or '').lower()
                roles.append((i, 'Primary' if not prim and i == ms[0] and 'table' not in t else ('Selectable metric' if 'selector' in t else ('Row / column in table' if 'table' in t or 'matrix' in t else 'Measure')), ''))
            for a in v.get('alt') or []:
                if a['metric'] and a['metric'] != prim:
                    roles.append((a['metric'], f'Alternate measure (switch: {a["label"]})', ''))
            for kid in v.get('kpis') or []:
                roles.append((kid, 'KPI card', ''))
        note = ''
        if generic:
            note = 'Generic layout: the metric is the one the user drilled from (any catalog metric).'
        if spec and not spec.get('metrics') and not v['metric']:
            note = (note + ' ' if note else '') + 'Metrics are not listed in the build annotation; see the prototype view.'
        if not spec:
            note = (note + ' ' if note else '') + 'No build annotation in the prototype; type inferred.'
        vtype = spec.get('type') or ('Detail table' if pid == 'accounts' else 'Table')
        dims = ', '.join(spec.get('dimensions') or [])
        sel = v.get('selects') or []
        if sel:
            dims = (dims + '; ' if dims else '') + 'Break down by: ' + ' / '.join(sel[0])
        rows_before = len(rows)
        add_visual(vid, title, vtype, dims, spec.get('interactions') or '', roles, note)
        if generic:
            rows[-1].update(metric=None, role='Selected metric (any catalog metric)', generic=True)
    if pid == 'definitions':
        for t in ['Client KPI coverage', 'Data sources', 'Dimensions', 'States', 'Data contract']:
            add_visual(vid, t, 'Reference table', '', 'Sub-tab', [], 'Reference content; see the matching tab in this workbook where one exists.')

# ------------------------------------------------------------------ workbook helpers
wb = Workbook()

def sheet(name, headers, data, widths, input_cols=(), table=True, wrap=True):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in data:
        ws.append(list(r))
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1] if c - 1 < len(widths) else 18
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=wrap, vertical='top')
            cell.border = Border(bottom=THIN)
            if headers[cell.column - 1] in input_cols:
                cell.fill = INPUT_FILL
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 32
    if table and data:
        ref = f'A1:{get_column_letter(len(headers))}{len(data) + 1}'
        t = Table(displayName='t' + ''.join(ch for ch in name if ch.isalnum()), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
        ws.add_table(t)
    return ws

def target_text(m):
    t = m['target']
    if t is None:
        return 'None'
    if t == 'dynamic':
        return f"Dynamic: cash goal = {CAT['cashGoalPctOfNpsr']:.0%} of lagged NPSR" if m['id'] in ('cash', 'cash_goal') else 'Dynamic'
    return fmt_val(t, m['unit'])

def fmt_val(v, unit):
    if v is None:
        return ''
    if isinstance(v, str):
        return v
    if unit == 'pct':
        return f'{v * 100:.4g}%'
    if unit == 'usd':
        return f'${v:,.0f}'
    return f'{v:g}'

VIEW_NAME = {v[0]: v[3] for v in views}
VIEW_DB = {v[0]: v[1] for v in views}
DB_NAME = {d[0]: d[1] for d in DASHBOARDS}
DB_DESC = {d[0]: d[2] for d in DASHBOARDS}
VIEW_DESC = {v[0]: f"{PAGE[v[2]]['question']} {VIEW_EXTRA.get(v[2], '')}".strip() for v in views}
VZ = {v[0]: v for v in visuals}

# ------------------------------------------------------------------ README
ws = wb.active
ws.title = 'README'
readme = [
    ('RCM Analytics: dashboard specification', ''),
    ('Purpose', 'The full specification of the dashboards, views, visuals, metrics and capabilities in the RCM Analytics prototype, for the development team. Scope: ELT Summary and Analytics. Worklists (v2) are out of scope.'),
    ('Source', 'Extracted from the prototype repository, not typed by hand: metric catalog (src/engine/metrics.ts), ELT framework (src/engine/framework.ts), targets (config/targets.config.json), capabilities (src/config/priorities.ts), client KPI inventory (src/config/kpiCoverage.ts), and the visual inventory read from the running prototype (build annotations). Regenerate with scripts/export-spec.ts and scripts/build_spec_workbook.py.'),
    ('Status', 'DRAFT. Metric definitions and targets in the prototype are drafts and illustrative. Confirm each with the client finance team (see Open Questions).'),
    ('', ''),
    ('Hierarchy', 'Dashboard (2)  ->  View (14)  ->  Visual  ->  Metric. One dashboard has many views; one view has many visuals; one visual shows one or more metrics. One metric can appear in many visuals.'),
    ('Keys', 'Dashboards: DB-##. Views: V-##. Visuals: <View ID>-VZ## (unique). Metrics: the catalog Metric ID (e.g. gross_ar_days), the same ID the prototype code uses. P-codes (P1-P13) are the client ELT framework codes.'),
    ('', ''),
    ('Tabs', ''),
    ('Spec', 'The main sheet (your requested columns). One row per visual x metric. Parent values repeat on every row so you can filter and sort. Definition and calculation come from the Metric Catalog.'),
    ('Dashboards', 'One row per dashboard.'),
    ('Views', 'One row per view (tab or drill-through page), with question answered, entry points and filters. Visual count is a live formula.'),
    ('Visuals', 'One row per visual: type, dimensions, interaction. Metric count is a live formula.'),
    ('Metric Catalog', 'One row per metric (single source for definition, calculation, unit, direction, target). "Used in visuals" is a live formula; 0 = in the catalog but not on any ELT / Analytics visual.'),
    ('Filters & Dimensions', 'Global and page-level filters, where they apply, and the rules for which measures ignore them.'),
    ('Capabilities', 'Cross-cutting product capabilities (navigation, filtering, every-visual tools, sharing, data trust), with scope.'),
    ('Data States', 'Non-happy-path states every visual must support.'),
    ('Client KPI Coverage', 'Every KPI on the client inventory sheet and how the product covers it (metric, filter/view, proxy, or source needed).'),
    ('Open Questions', 'Assumptions and decisions to confirm before build.'),
    ('Glossary', 'Terms used in this workbook.'),
    ('', ''),
    ('Legend', 'Yellow cells = to be filled in (Time Recency, Release, Owner). Release uses a drop-down: v1 (MVP / Must), v2 (Should), v3+ (Could), Won\'t.'),
    ('', ''),
    ('Changes to the requested column list', ''),
    ('1. Visual level added', 'A view holds many visuals and the same metric shows in many of them (e.g. Cash % NPSR as a KPI card, a hospital bar chart and a trend). "View Metrics Visualization" is therefore a visual, with its own ID.'),
    ('2. Metric Catalog tab', 'Definition and calculation live once in the catalog; Spec repeats them for reading. Edit the catalog, then regenerate, so the two never drift.'),
    ('3. Stable IDs', 'IDs at every level so tabs join and the team can trace a row to code and to tickets.'),
    ('4. Extra spec columns', 'Metric role in the visual, dimensions, interaction / drill target, unit, direction, flow vs balance, target, release. These are what a developer needs to build the visual, beyond its name.'),
    ('5. Time Recency', 'Left blank as asked. Suggest splitting it into: data grain (day / month), refresh cadence (e.g. daily), data latency (as-of lag), and period basis (flow = sum over period; balance = value at period end). Period basis is already filled from the catalog.'),
    ('6. MECE', 'Spec rows are mutually exclusive (one per visual x metric). Completeness checks: every catalog metric is either used (Used in visuals > 0) or flagged; every client KPI has a coverage row; capabilities and states are listed once and referenced by scope, not repeated per visual.'),
    ('7. Not in this workbook', 'Worklists (v2), prototype-only controls (prototype bar, release prioritization, demo scenarios). Requirements (BR / FR / TR) stay in the requirements workbook; add a Trace column there that points to these IDs.'),
]
for r in readme:
    ws.append(list(r))
ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 130
for row in ws.iter_rows():
    for c in row:
        c.font = Font(name=FONT, size=10, bold=(c.column == 1))
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(name=FONT, size=14, bold=True)
for r in (9, 25):
    ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True, color='1F3864')

# ------------------------------------------------------------------ Dashboards / Views / Visuals
sheet('Dashboards', ['Dashboard ID', 'Dashboard', 'Dashboard Description', 'Primary users', 'Navigation', 'Filter model', 'Owner'],
      DASHBOARDS + [], [12, 16, 70, 30, 40, 40, 16], input_cols=('Owner',))
for i in range(2, 4):
    wb['Dashboards'].cell(row=i, column=7).fill = INPUT_FILL

vrows = []
for k, (vid, db, pid, name, vtype) in enumerate(views):
    r = k + 2
    vrows.append((vid, db, DB_NAME[db], name, VIEW_DESC[vid], PAGE[pid]['question'], vtype, ENTRY.get(pid, 'Tab' if db == 'DB-02' else 'Dashboard home'),
                  page_filters(pid), f'=COUNTIF(Visuals!$B:$B,A{r})', ''))
sheet('Views', ['View ID', 'Dashboard ID', 'Dashboard', 'View', 'View Description', 'Question answered', 'View type', 'Entry points', 'Filters available', 'Visual count', 'Release'],
      vrows, [9, 11, 13, 30, 70, 45, 13, 32, 40, 9, 11], input_cols=('Release',))

zrows = []
for k, (vz, vid, title, vtype, dims, inter, notes) in enumerate(visuals):
    r = k + 2
    zrows.append((vz, vid, VIEW_NAME[vid], title, vtype, dims, inter, f'=COUNTIF(Spec!$G:$G,A{r})', '', notes))
sheet('Visuals', ['Visual ID', 'View ID', 'View', 'Visual', 'Visual type', 'Dimensions', 'Interaction', 'Metric rows', 'Release', 'Notes'],
      zrows, [14, 8, 26, 40, 40, 30, 50, 9, 11, 40], input_cols=('Release',))

# ------------------------------------------------------------------ Spec
srows = []
for k, r in enumerate(rows):
    vz = VZ[r['vz']]
    vid = vz[1]
    db = VIEW_DB[vid]
    m = M.get(r['metric']) if r['metric'] else None
    srows.append((
        f'S-{k + 1:04d}', DB_NAME[db], DB_DESC[db], VIEW_NAME[vid], VIEW_DESC[vid],
        vid, vz[0], vz[2],
        mname(m['id']) if m else ('{Selected metric}' if r.get('generic') else ''),
        m['id'] if m else '', m['code'] if m else '',
        r['role'],
        (m['calculation'] if m else ('See Metric Catalog for the selected metric' if r.get('generic') else '')),
        '',  # Time recency (blank by request)
        vz[3],
        vz[4], vz[5],
        m['definition'] if m else '',
        UNIT[m['unit']] if m else '', DIRN[m['direction']] if m else '',
        ('Flow (sum over period)' if m['type'] == 'flow' else 'Balance (at period end)') if m else '',
        target_text(m) if m else '', '', r.get('notes') or vz[6],
    ))
SPEC_HEAD = ['Row ID', 'Dashboard', 'Dashboard Description', 'View', 'View Description', 'View ID', 'Visual ID', 'Visual',
             'View Metrics', 'Metric ID', 'P-code', 'Metric role in visual', 'View Metrics Calculation', 'View Metrics Time Recency',
             'View Metrics Visualization', 'Dimensions', 'Interaction / drill', 'Metric definition', 'Unit', 'Direction', 'Period basis',
             'Target', 'Release', 'Notes']
ws = sheet('Spec', SPEC_HEAD, srows, [9, 13, 40, 24, 45, 8, 13, 34, 32, 24, 7, 22, 55, 14, 34, 24, 40, 50, 9, 14, 16, 14, 11, 35],
           input_cols=('View Metrics Time Recency', 'Release'))
ws.freeze_panes = 'J2'
wb.move_sheet('Spec', offset=-(len(wb.sheetnames) - 2))

# ------------------------------------------------------------------ Metric Catalog
crows = []
for k, m in enumerate(CAT['metrics']):
    r = k + 2
    grp = FW.get(m['id'])
    crows.append((m['id'], m['code'], grp[0] if grp else '', m['name'], m['short'], m['area'], m['definition'], m['calculation'],
                  UNIT[m['unit']], DIRN[m['direction']], 'Flow' if m['type'] == 'flow' else 'Balance', ', '.join(m['sourceFields']),
                  'Yes' if m['multiDef'] else 'No', target_text(m), fmt_val(m['watch'], m['unit']) if m['target'] != 'dynamic' else (f"{m['watch']:.0%} of goal" if isinstance(m['watch'], (int, float)) else ''),
                  grp[1]['benchmark'] if grp else '', PAGE['executive' if m['page'] == 'executive' else m['page']]['title'] if (m['page'] in PAGE) else m['page'],
                  f'=COUNTIF(Spec!$J:$J,A{r})', ''))
sheet('Metric Catalog', ['Metric ID', 'P-code', 'Framework group', 'Metric', 'Short name', 'RCM area', 'Definition', 'Calculation', 'Unit', 'Direction',
                         'Flow / Balance', 'Source fields', 'More than one industry definition', 'Target', 'Watch threshold', 'Benchmark (client)',
                         'Primary analysis view', 'Used in visuals', 'Time Recency'],
      crows, [26, 7, 20, 34, 14, 16, 55, 55, 9, 18, 9, 40, 12, 14, 12, 16, 22, 9, 14], input_cols=('Time Recency',))

# ------------------------------------------------------------------ Filters
an_views = ', '.join(VIEW_BY_PAGE[p][0] for p in ANALYTICS_ORDER)
frows = [
    ('F-01', 'ELT slicer', 'Period type', 'Month, Quarter, Year to date, Rolling 12 months', 'V-01', 'All metrics', ''),
    ('F-02', 'ELT slicer', 'Period', 'Periods available in the data window', 'V-01', 'All metrics', ''),
    ('F-03', 'ELT slicer', 'Hospital (or region)', 'All hospitals, one hospital, or multiple selected', 'V-01', 'All metrics', 'Other active filters show as a chip with a clear button.'),
    ('F-04', 'Global (Date)', 'Period type', 'Month, Quarter, Year to date, Rolling 12 months', an_views, 'All metrics', ''),
    ('F-05', 'Global (Date)', 'Period', 'Periods available for the period type', an_views, 'All metrics', 'Shows start and end date.'),
    ('F-06', 'Global (Date)', 'Compare to', 'Prior period, Prior year', an_views, 'Change and comparison columns', ''),
]
n = 7
for sec, fields in [('Global (Organization)', ['region', 'facilityType', 'facility']), ('Global (Payer)', ['financialClass', 'payer']), ('Global (Service)', ['patientType', 'serviceLine'])]:
    for f in fields:
        frows.append((f'F-{n:02d}', sec, FIELD[f], 'Multi-select with search, select all, clear', an_views,
                      'Metrics whose source carries this field; others show "measure not available by this filter"', ''))
        n += 1
frows.append((f'F-{n:02d}', 'Global (Metrics)', 'RCM area', 'Patient Access, Charge Capture, Coding, Clinical Documentation, Billing, Denials, A/R, Cash, Financial', an_views, 'Scorecards (limits metric rows)', '')); n += 1
NOTE = {'arAge': 'Applies to A/R balances ($). A/R days and aging shares ignore it.', 'accountStatus': 'Applies to A/R balances ($). A/R days and aging shares ignore it.',
        'denialCategory': 'Applies to denial measures.', 'rootCause': 'Applies to denial measures.', 'dnfbHold': 'Applies to DNFB measures.', 'editCategory': 'Applies to claim edit measures.'}
for pid in ANALYTICS_ORDER:
    for f in PAGE[pid].get('pageFilters') or []:
        frows.append((f'F-{n:02d}', 'Page filter', FIELD[f], 'Multi-select', VIEW_BY_PAGE[pid][0] + ' ' + PAGE[pid]['title'], NOTE[f], '')); n += 1
frows.append((f'F-{n:02d}', 'Interaction', 'Click-to-filter (cross-filter)', 'Click a bar, segment, row or point', an_views, 'Every visual on the view', 'Unselected members fade. Applied filters show as removable chips; Clear all.')); n += 1
frows.append((f'F-{n:02d}', 'Behavior', 'Filter persistence', '', 'All', '', 'Filters persist between views and visits (prototype: browser storage).'))
sheet('Filters & Dimensions', ['Filter ID', 'Level', 'Filter', 'Values / control', 'Views', 'Applies to', 'Notes'], frows, [9, 20, 26, 40, 40, 50, 45])

# ------------------------------------------------------------------ Capabilities
caprows = []
for k, c in enumerate(CAT['capabilities']):
    out = c['group'] == 'Worklists' or c['id'] == 'cap:see-understand-act'
    scope = 'Out of scope (Worklists v2)' if out else ('Every visual' if c['group'] == 'Every visual' else ('ELT Summary + Analytics' if c['group'] in ('Data trust', 'Sharing', 'Users and access', 'Navigation') else 'Analytics'))
    caprows.append((f'C-{k + 1:02d}', c['group'], c['label'], scope, ''))
sheet('Capabilities', ['Capability ID', 'Group', 'Capability', 'Scope', 'Release'], caprows, [12, 18, 70, 30, 11], input_cols=('Release',))

# ------------------------------------------------------------------ Data states
st = [
    ('DS-01', 'Loading', 'Data request in progress', 'Skeleton placeholder in each visual'),
    ('DS-02', 'No data for filters', 'Filters return no rows', 'Plain message "No data matches the current filters" with a Clear filters button'),
    ('DS-03', 'Measure not available by a filter', 'A filter is active that the metric\'s source does not carry (e.g. payer on a call-center metric)', 'Message naming the filter that the measure cannot be split by'),
    ('DS-04', 'Source unavailable', 'A source feed (e.g. clearinghouse) has not loaded', 'Message naming the source and the last load time; other visuals unaffected'),
    ('DS-05', 'Error with retry', 'Data service returns an error', 'Error message with request id and a Retry button; other visuals unaffected'),
    ('DS-06', 'Stale data', 'Last refresh older than the expected schedule', 'Banner at the top of the view'),
    ('DS-07', 'Preliminary period', 'Selected period is not closed', '"Preliminary" label in the header next to the reporting period'),
]
sheet('Data States', ['State ID', 'State', 'Trigger', 'Expected display'], st, [9, 30, 55, 70])

# ------------------------------------------------------------------ Client KPI coverage
krows = [(f'K-{k + 1:02d}', x['dashboard'], x['stage'], x['kpi'], x['coverage'], ', '.join(x['metrics']), x['how'], x.get('source') or '', '')
         for k, x in enumerate(CAT['kpiCoverage'])]
sheet('Client KPI Coverage', ['KPI ID', 'Client dashboard', 'Stage', 'Client KPI', 'Coverage', 'Metric IDs', 'Where / how', 'Data still needed', 'Release'],
      krows, [8, 24, 11, 40, 16, 34, 55, 50, 11], input_cols=('Release',))

# ------------------------------------------------------------------ Open questions
unused = [m['id'] for m in CAT['metrics'] if not any(r['metric'] == m['id'] for r in rows)]
multi = [m['id'] for m in CAT['metrics'] if m['multiDef']]
oq = [
    ('Metric definitions', 'All catalog definitions are drafts. ' + str(len(multi)) + ' metrics have more than one common industry definition (column "More than one industry definition" = Yes). Confirm each with the client finance team.'),
    ('Targets', 'Targets and watch thresholds are illustrative prototype configuration (config/targets.config.json). ELT watch thresholds are prototype choices, not client values.'),
    ('P5 DNFB target', 'Target = bill hold days + 1.5. Bill hold days = 3 in the prototype (illustrative), so the P5 target is 4.5 days. Confirm bill hold per client.'),
    ('P6 Credit balance days', 'Credit balances are modeled; the synthetic data has no overpayments. Confirm the source for credit balance AR.'),
    ('P8 Cash as % of net revenue', 'Uses average monthly net revenue over the most recent 3 months, including the reporting month; quarters and YTD use the monthly average x months. Confirm.'),
    ('P9 Bad debt', 'The provision is the bad debt transfers posted in the period. Confirm against the GL provision.'),
    ('P10 Initial denial rate', 'Denials by denial date over claims by submit date in the same period: an approximation of the submission cohort. Confirm cohort vs period basis.'),
    ('P11 Unrealized bad debt', 'Unrealized = open self-pay balances older than 120 days after discharge. Confirm the criteria.'),
    ('P12 Avoidable write-offs', 'Avoidable write-offs = write-offs after a final denial. Confirm the write-off codes in scope.'),
    ('P13 Unrealized avoidable', 'Unrealized = expected net on open denials older than 90 days. Confirm.'),
    ('Account-level detail', 'Account Detail (V-14) and "Largest open accounts" (A/R) list account-level records. If this product is client-facing, check this against requirement BR-32 / FR-33 (no account-level records in v1).'),
] + ([] if not unused else [
    ('Catalog metrics not on any visual', f'{len(unused)} catalog metrics are not on any ELT / Analytics visual (Used in visuals = 0): ' + ', '.join(unused) + '. Some are reached only through Metric Analysis, the Client KPI Coverage list or worklists. Decide: show, keep as drill-only, or drop.'),
]) + [
    ('Mid-Cycle view defect', 'In the prototype dev build the Mid-Cycle view throws a chart library error (ECharts axis assertion) and does not render. The spec rows were read with a local guard. Fix in the prototype.'),
    ('Generic drill-through', 'Metric Analysis (V-12) uses one layout for every metric. Confirm that each metric supports every breakdown (hospital, payer) or which breakdowns to hide.'),
    ('Time Recency', 'Not yet defined. Needed per metric: data grain, refresh cadence, latency, and the close calendar.'),
    ('Internal vs External users', 'The prototype has an Internal / External user switch that shows the same view today. Decide which views, metrics and capabilities differ by user type.'),
]
sheet('Open Questions', ['Topic', 'Question / assumption', 'Owner', 'Decision', 'Date'], oq, [28, 110, 16, 40, 12], input_cols=('Owner', 'Decision', 'Date'))

# ------------------------------------------------------------------ Glossary
gl = [
    ('Dashboard', 'A top-level experience with its own users and navigation (ELT Summary, Analytics).'),
    ('View', 'One page of a dashboard: a tab or a drill-through page.'),
    ('Visual', 'One chart, table, KPI card group or panel on a view.'),
    ('Metric', 'One calculated measure from the Metric Catalog.'),
    ('Drill-through', 'Opening a more detailed view from a visual, carrying the current filters.'),
    ('Flow metric', 'Summed over the period (e.g. cash collected in the month).'),
    ('Balance metric', 'Value at the period end snapshot (e.g. A/R at month end).'),
    ('Ratio rule', 'Ratios are sum(numerator) / sum(denominator) for the filter context, never an average of ratios.'),
    ('Target / Watch', 'Target = goal. Watch = the limit between Watch and Off target. Status = On target, Watch, or Off target.'),
    ('Direction', 'Whether a higher or lower value is favorable; controls variance colors.'),
    ('Preliminary period', 'A period that is not closed; values can still change.'),
    ('P-code', 'Code in the client ELT metric framework (P1-P13; there is no P3).'),
    ('NPSR', 'Net patient service revenue (catalog metric npsr).'),
    ('DNFB', 'Discharged not final billed: accounts discharged but not yet final billed.'),
    ('Billed, not exported', 'Claims final billed but not yet sent to the payer, usually held by open claim edits (prototype measures dnsp_*).'),
    ('RCM area', 'Revenue cycle stage: Patient Access, Charge Capture, Coding, Clinical Documentation, Billing, Denials, A/R, Cash, Financial.'),
]
sheet('Glossary', ['Term', 'Meaning'], gl, [26, 110])

# Release drop-downs
for name, col in [('Spec', 'W'), ('Views', 'K'), ('Visuals', 'I'), ('Capabilities', 'E'), ('Client KPI Coverage', 'I')]:
    w = wb[name]
    dv = DataValidation(type='list', formula1='"v1,v2,v3+,Won\'t"', allow_blank=True)
    w.add_data_validation(dv)
    dv.add(f'{col}2:{col}{w.max_row}')

wb.save(OUT)
print(f'views {len(views)}, visuals {len(visuals)}, spec rows {len(rows)}, metrics {len(CAT["metrics"])}, unused {len(unused)}')
